"""Generates a complete, ready-to-present PowerPoint slide deck for DocAssistIQ.
Includes the high-res architecture diagram and key architectural takeaways.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # -----------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -----------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Background shape
    bg = slide1.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5)) # MSO_SHAPE.RECTANGLE = 1
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 41, 66) # Deep Navy
    bg.line.color.rgb = RGBColor(15, 41, 66)

    # Title box
    txBox = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "DocAssistIQ"
    p.font.name = "Calibri"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = "System Architecture & Clinical Decision Support Engineering"
    p2.font.name = "Calibri"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
    p2.space_before = Pt(12)
    
    p3 = tf.add_paragraph()
    p3.text = "Neuro-Symbolic Architecture • Real-Time Ambient Voice • Deterministic Safety • Local GPU Inference"
    p3.font.name = "Calibri"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184) # Slate
    p3.space_before = Pt(18)

    # -----------------------------------------------------------------------
    # SLIDE 2: Primary Architecture Slide (16:9 High-Res Diagram)
    # -----------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    img_path = os.path.abspath("DocAssistIQ_Architecture_PPT_Slide.png")
    if os.path.exists(img_path):
        # Center the diagram on the slide with slight padding
        slide2.shapes.add_picture(img_path, Inches(0.4), Inches(0.3), width=Inches(12.533))

    # -----------------------------------------------------------------------
    # SLIDE 3: Key Architectural Highlights
    # -----------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    
    # Slide Title
    tx = slide3.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.8))
    p = tx.text_frame.paragraphs[0]
    p.text = "DocAssistIQ: Architectural Pillars & Design Innovations"
    p.font.name = "Calibri"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 41, 66)

    # 4 Pillar cards
    pillars = [
        ("1. Ambient Voice Pipeline", "• Web Audio API (16kHz PCM)\n• Faster-Whisper Large-v3\n• PyAnnote 3.1 Diarization\n• Mandatory Patient Consent Gate", RGBColor(2, 132, 199)),
        ("2. Deterministic Reasoning", "• 100+ Multi-Specialty Disease KB\n• Multi-Factor Jaccard + Cardinal Math\n• Incubation & Travel Fit\n• <85ms Sub-Second Execution", RGBColor(13, 148, 136)),
        ("3. Zero-Bypass Safety", "• Master Red Flags (RF-001 to 005)\n• 100% Sensitivity on Emergencies\n• Live NIH RxNav DDI Screening\n• Mandatory Clinician Review", RGBColor(185, 28, 28)),
        ("4. Local Privacy & Security", "• Constrained Llama-3.2-3B via Ollama\n• PostgreSQL 16 + pgvector HNSW\n• SHA-256 Chained Audit Trail\n• Zero Cloud Data Leakage (HIPAA)", RGBColor(124, 58, 237)),
    ]

    card_w = Inches(2.7)
    card_h = Inches(4.8)
    y_pos = Inches(1.8)

    for i, (ptitle, ptext, pcol) in enumerate(pillars):
        x_pos = Inches(0.8 + (i * 2.95))
        box = slide3.shapes.add_shape(1, x_pos, y_pos, card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(248, 250, 252)
        box.line.color.rgb = pcol
        box.line.width = Pt(2)

        tf = box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = ptitle
        p.font.name = "Calibri"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = pcol
        
        p2 = tf.add_paragraph()
        p2.text = ptext
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = RGBColor(30, 41, 59)
        p2.space_before = Pt(14)

    output_pptx = os.path.abspath("DocAssistIQ_Clean_Headings_Presentation.pptx")
    prs.save(output_pptx)
    print(f"Generated PowerPoint Presentation: {output_pptx}")
    
    # Try updating the default name if unlocked
    try:
        orig_pptx = os.path.abspath("DocAssistIQ_Architecture_Presentation.pptx")
        prs.save(orig_pptx)
        print(f"Also updated: {orig_pptx}")
    except PermissionError:
        print("Note: Original DocAssistIQ_Architecture_Presentation.pptx is open in PowerPoint. Saved to DocAssistIQ_Clean_Headings_Presentation.pptx instead.")
        
    return output_pptx

if __name__ == "__main__":
    create_presentation()
